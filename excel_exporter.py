"""
excel_exporter.py
=================
Escribe el reporte final en un archivo Excel con varias hojas y formato:
- Portada con KPIs y resumen de calidad de datos
- Tablas de análisis
- Gráficas incrustadas
"""

from __future__ import annotations

from pathlib import Path
from typing import Mapping

import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import column_index_from_string
from openpyxl.worksheet.worksheet import Worksheet


# ---------------------------------------------------------------------------
# Estilos compartidos
# ---------------------------------------------------------------------------

_TITLE_FILL = PatternFill("solid", fgColor="1F2937")
_TITLE_FONT = Font(name="Calibri", size=18, bold=True, color="FFFFFF")
_SECTION_FILL = PatternFill("solid", fgColor="2563EB")
_SECTION_FONT = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="E5E7EB")
_HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="111827")
_TOTAL_FILL = PatternFill("solid", fgColor="FEF3C7")
_TOTAL_FONT = Font(name="Calibri", size=11, bold=True, color="92400E")
_OK_FONT = Font(color="065F46", bold=True)
_BAD_FONT = Font(color="991B1B", bold=True)

_THIN = Side(border_style="thin", color="9CA3AF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _autosize(ws: Worksheet, max_width: int = 38, min_width: int = 10) -> None:
    for col_cells in ws.columns:
        # Cada tupla de openpyxl incluye la celda vacía en algunos extremos;
        # filtramos para quedarnos con valores reales.
        valid = [c for c in col_cells if c.value is not None]
        if not valid:
            continue
        col_letter = get_column_letter(valid[0].column)
        longest = 0
        for cell in valid:
            length = len(str(cell.value))
            if length > longest:
                longest = length
        width = max(min_width, min(max_width, longest + 2))
        ws.column_dimensions[col_letter].width = width


def _write_title(ws: Worksheet, text: str, span: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
    cell = ws.cell(row=1, column=1, value=text)
    cell.fill = _TITLE_FILL
    cell.font = _TITLE_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32


def _write_section(ws: Worksheet, row: int, text: str, span: int) -> None:
    ws.merge_cells(
        start_row=row, start_column=1, end_row=row, end_column=span
    )
    cell = ws.cell(row=row, column=1, value=text)
    cell.fill = _SECTION_FILL
    cell.font = _SECTION_FONT
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row].height = 22


def _write_table(
    ws: Worksheet,
    df: pd.DataFrame,
    start_row: int,
    number_columns: tuple[str, ...] = (),
) -> int:
    """Escribe un DataFrame con cabecera. Devuelve la siguiente fila libre."""
    headers = list(df.columns)
    n_cols = len(headers)
    header_row = start_row

    for col_idx, name in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=col_idx, value=name)
        c.fill = _HEADER_FILL
        c.font = _HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = _BORDER

    for r_idx, row in enumerate(df.itertuples(index=False), start=header_row + 1):
        is_total = str(row[0]).strip().lower() in {"totales", "total"}
        for c_idx, value in enumerate(row, start=1):
            c = ws.cell(row=r_idx, column=c_idx, value=value)
            c.border = _BORDER
            c.alignment = Alignment(vertical="center")
            col_name = headers[c_idx - 1]
            if col_name in number_columns and isinstance(value, (int, float)):
                c.number_format = "#,##0.00"
                c.alignment = Alignment(horizontal="right", vertical="center")
            if is_total:
                c.fill = _TOTAL_FILL
                c.font = _TOTAL_FONT
    return header_row + len(df) + 1


# ---------------------------------------------------------------------------
# Hojas
# ---------------------------------------------------------------------------


def _build_kpis_sheet(
    wb: Workbook, kpis: pd.DataFrame, summary, rejected: pd.DataFrame
) -> None:
    ws = wb.active
    ws.title = "Resumen"
    _write_title(ws, "Reporte Automatizado de Ventas", 4)

    _write_section(ws, 3, "Indicadores generales", 4)
    next_row = _write_table(ws, kpis, 4)

    _write_section(ws, next_row, "Calidad de los datos", 4)
    quality = (
        pd.DataFrame(
            {
                "Métrica": [
                    "Filas leídas",
                    "Filas aceptadas",
                    "Filas rechazadas",
                    "% Aceptadas",
                ],
                "Valor": [
                    summary.total_rows,
                    summary.accepted_rows,
                    summary.rejected_rows,
                    f"{(summary.accepted_rows / max(summary.total_rows, 1) * 100):.2f}%",
                ],
            }
        )
        if summary.total_rows
        else pd.DataFrame({"Métrica": [], "Valor": []})
    )
    next_row = _write_table(ws, quality, next_row + 1, number_columns=())

    if not rejected.empty:
        _write_section(ws, next_row, "Filas descartadas (motivos)", 4)
        rej_cols = [
            "row_number",
            "order_id",
            "date",
            "product",
            "quantity",
            "unit_price",
            "customer_email",
            "motivos",
        ]
        rej_df = rejected[rej_cols].copy()
        rej_df.columns = [
            "Fila",
            "Order ID",
            "Fecha",
            "Producto",
            "Cantidad",
            "Precio",
            "Email",
            "Motivos",
        ]
        next_row = _write_table(ws, rej_df, next_row + 1)
        # Coloreamos en rojo la columna Motivos para reforzar visualmente.
        for row in ws.iter_rows(
            min_row=next_row - len(rej_df),
            max_row=next_row - 1,
            min_col=1,
            max_col=8,
        ):
            motivo_cell = row[-1]
            motivo_cell.font = _BAD_FONT
    else:
        ws.cell(row=next_row, column=1, value="No se rechazó ninguna fila.").font = _OK_FONT

    _autosize(ws)
    ws.column_dimensions["A"].width = 22


def _build_analysis_sheets(
    wb: Workbook,
    sheets: Mapping[str, pd.DataFrame],
) -> None:
    for name, df in sheets.items():
        ws = wb.create_sheet(title=name[:31])
        _write_title(ws, name, len(df.columns))
        _write_table(ws, df, 3)
        _autosize(ws)


def _build_charts_sheet(wb: Workbook, charts: Mapping[str, Path]) -> None:
    ws = wb.create_sheet("Gráficas")
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 90
    _write_title(ws, "Gráficas", 2)
    row = 3
    for title, path in charts.items():
        if path is None or not path.exists():
            continue
        _write_section(ws, row, title, 2)
        ws.row_dimensions[row + 1].height = 300
        try:
            img = XLImage(str(path))
            img.width = 900
            img.height = 480
            ws.add_image(img, f"B{row + 1}")
            row += 22  # deja espacio hasta la siguiente sección
        except Exception:
            # Si falla la imagen, dejamos el espacio en blanco pero visible.
            ws.cell(row=row + 1, column=2, value=f"[No se pudo incrustar {path.name}]")
            row += 8


def export_excel(
    output_path: Path,
    kpis: pd.DataFrame,
    summary,
    rejected: pd.DataFrame,
    analysis_sheets: Mapping[str, pd.DataFrame],
    charts: Mapping[str, Path],
) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    _build_kpis_sheet(wb, kpis, summary, rejected)
    _build_analysis_sheets(wb, analysis_sheets)
    _build_charts_sheet(wb, charts)
    wb.save(output_path)
    return output_path


__all__ = ["export_excel"]
